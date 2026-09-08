from __future__ import annotations

import re
from collections.abc import Iterable
from datetime import date, datetime, time as dt_time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import (
	Column,
	MetaData,
	Table,
	Text,
	create_engine,
	select,
	inspect,
	text,
	Integer,
	Float,
	Boolean,
)
from sqlalchemy.types import DateTime as SA_DateTime, Date as SA_Date, Time as SA_Time

DATABASE_URL_EXAMPLES = {
	'SQLite': 'sqlite:///files_to_sql.db',
	'PostgreSQL': 'postgresql+psycopg://usuario@localhost:5432/meu_banco',
	'MySQL': 'mysql+pymysql://usuario@localhost:3306/meu_banco',
	'SQL Server': 'mssql+pyodbc://usuario@localhost/meu_banco?driver=ODBC+Driver+18+for+SQL+Server',
}

DATABASE_URL = 'mysql+pymysql://root:admin@localhost:3306/idl'
# Configure os valores aqui no topo em vez de usar argumentos de linha de comando
TABLES_DIR = Path('tabelas')


def normalize_table_name(file_path: Path) -> str:
	table_name = re.sub(r'\W+', '_', file_path.stem.strip().lower()).strip('_')
	if not table_name:
		raise ValueError(f'Não foi possível gerar um nome de tabela para {file_path.name}.')
	return table_name


def normalize_column_names(headers: Iterable[Any]) -> list[str]:
	"""Normalize headers to lowercase, replace non-word with underscore and ensure uniqueness."""
	normalized: list[str] = []
	seen: dict[str, int] = {}

	for index, header in enumerate(headers, start=1):
		value = str(header).strip() if header not in (None, '') else f'coluna_{index}'
		if not value:
			value = f'coluna_{index}'

		value = re.sub(r'\W+', '_', value).strip('_').lower()
		if not value:
			value = f'coluna_{index}'

		count = seen.get(value, 0)
		seen[value] = count + 1
		normalized.append(value if count == 0 else f'{value}_{count + 1}')

	return normalized


def normalize_cell_value(value: Any) -> Any:
	"""Keep original cell values (do not coerce/format)."""
	return value


def infer_single_column_type(values: list[Any]):
	"""Infer SQLAlchemy column type from sample values."""
	encountered_int = encountered_float = encountered_bool = False
	encountered_datetime = encountered_date = encountered_time = encountered_str = False

	for v in values:
		if v is None or (isinstance(v, str) and v.strip() == ''):
			continue

		if isinstance(v, bool):
			encountered_bool = True
			continue

		if isinstance(v, int) and not isinstance(v, bool):
			encountered_int = True
			continue

		if isinstance(v, float):
			encountered_float = True
			continue

		if isinstance(v, datetime):
			encountered_datetime = True
			continue

		if isinstance(v, date) and not isinstance(v, datetime):
			encountered_date = True
			continue

		if isinstance(v, dt_time):
			encountered_time = True
			continue

		if isinstance(v, str):
			s = v.strip()
			low = s.lower()
			if low in ('true', 'false', 't', 'f', 'yes', 'no', 'y', 'n', 'sim', 'nao', 'não'):
				encountered_bool = True
				continue

			if re.fullmatch(r'[+-]?\d+', s):
				encountered_int = True
				continue

			if re.fullmatch(r'[+-]?(?:\d*\.\d+|\d+\.\d*)(?:[eE][+-]?\d+)?', s) or re.fullmatch(
				r'[+-]?\d+[eE][+-]?\d+', s
			):
				encountered_float = True
				continue

			# try iso parsing for dates/times
			try:
				datetime.fromisoformat(s)
				encountered_datetime = True
				continue
			except Exception:
				pass

			try:
				date.fromisoformat(s)
				encountered_date = True
				continue
			except Exception:
				pass

			try:
				dt_time.fromisoformat(s)
				encountered_time = True
				continue
			except Exception:
				pass

			encountered_str = True

	# Prefer concrete date/time types only when there are no non-date/time strings
	# in the sample. If the column contains a mix of date/time objects and
	# arbitrary strings (common with Excel import), default to Text to avoid
	# binding errors (e.g. MySQL rejecting non-time strings for TIME columns).
	if encountered_datetime and not encountered_str:
		return SA_DateTime()
	if encountered_date and not encountered_str:
		return SA_Date()
	if encountered_time and not encountered_str:
		return SA_Time()
	if encountered_bool and not (encountered_int or encountered_float or encountered_str):
		return Boolean()
	if encountered_int and not (encountered_float or encountered_str):
		return Integer()
	if (encountered_float and not encountered_str) or (
		encountered_int and encountered_float and not encountered_str
	):
		return Float()

	return Text()


def infer_column_types(
	columns: list[str], rows: list[tuple[Any, ...]], sample_size: int = 200
) -> dict[str, Any]:
	samples: dict[int, list[Any]] = {i: [] for i in range(len(columns))}
	for row in rows[:sample_size]:
		for i in range(len(columns)):
			val = row[i] if i < len(row) else None
			samples[i].append(val)

	result: dict[str, Any] = {}
	for i, col in enumerate(columns):
		result[col] = infer_single_column_type(samples[i])
	return result


def cast_value_to_type(value: Any, sa_type: Any) -> Any:
	"""Attempt a safe cast of a value to the inferred SQLAlchemy type when reasonable."""
	if value is None or value == '':
		return None

	try:
		if isinstance(sa_type, Integer):
			if isinstance(value, int) and not isinstance(value, bool):
				return int(value)
			if isinstance(value, float) and value.is_integer():
				return int(value)
			if isinstance(value, str) and re.fullmatch(r'[+-]?\d+', value.strip()):
				return int(value.strip())
			return value

		if isinstance(sa_type, Float):
			if isinstance(value, (int, float)):
				return float(value)
			if isinstance(value, str) and re.fullmatch(
				r'[+-]?(?:\d*\.\d+|\d+\.\d*)(?:[eE][+-]?\d+)?', value.strip()
			):
				return float(value.strip())
			return value

		if isinstance(sa_type, Boolean):
			if isinstance(value, bool):
				return value
			if isinstance(value, (int, float)):
				return bool(value)
			if isinstance(value, str):
				low = value.strip().lower()
				if low in ('true', 't', 'yes', 'y', 'sim'):
					return True
				if low in ('false', 'f', 'no', 'n', 'nao', 'não'):
					return False
			return value

		if isinstance(sa_type, (SA_DateTime, SA_Date, SA_Time)):
			# If value is already a date/time object, keep it.
			if isinstance(value, (datetime, date, dt_time)):
				return value
			if isinstance(value, str):
				s = value.strip()
				try:
					if isinstance(sa_type, SA_DateTime):
						return datetime.fromisoformat(s)
					if isinstance(sa_type, SA_Date):
						return date.fromisoformat(s)
					if isinstance(sa_type, SA_Time):
						return dt_time.fromisoformat(s)
				except Exception:
					return value

		return value
	except Exception:
		return value


def create_engine_from_url(database_url: str):
	return create_engine(database_url)


def unique_column_name(base: str, used: set[str]) -> str:
	name = base
	index = 2
	while name in used:
		name = f'{base}_{index}'
		index += 1
	used.add(name)
	return name


def is_id_name(column_name: str) -> bool:
	return column_name.strip().lower() == 'id'


def remap_incoming_columns(columns: list[str]) -> list[str]:
	"""Map incoming columns to DB-safe names, reserving `id` for technical PK."""
	used = {'id'}
	result: list[str] = []

	for column in columns:
		base = 'source_id' if is_id_name(column) else column
		result.append(unique_column_name(base, used))

	return result


def make_temp_table_name(table_name: str, max_len: int = 60) -> str:
	temp = f'__tmp_{table_name}_idpk'
	if len(temp) <= max_len:
		return temp

	trim = max_len - len('__tmp__idpk')
	trimmed = table_name[: max(trim, 1)]
	return f'__tmp_{trimmed}_idpk'


def rename_table(connection, dialect_name: str, from_name: str, to_name: str) -> None:
	preparer = connection.dialect.identifier_preparer
	quoted_from = preparer.quote(from_name)
	quoted_to = preparer.quote(to_name)

	if dialect_name == 'mysql':
		connection.execute(text(f'RENAME TABLE {quoted_from} TO {quoted_to}'))
		return

	if dialect_name in {'mssql'}:
		connection.execute(text(f"EXEC sp_rename '{from_name}', '{to_name}'"))
		return

	connection.execute(text(f'ALTER TABLE {quoted_from} RENAME TO {quoted_to}'))


def table_has_id_primary_key(engine, table_name: str) -> bool:
	inspector = inspect(engine)
	pk = inspector.get_pk_constraint(table_name) or {}
	pk_columns = {str(col).lower() for col in (pk.get('constrained_columns') or [])}
	columns = {str(col['name']).lower() for col in inspector.get_columns(table_name)}
	return 'id' in columns and pk_columns == {'id'}


def rebuild_table_with_id_primary_key(engine, table_name: str) -> None:
	"""Rebuild an existing table to ensure `id` is the only primary key."""
	metadata = MetaData()
	existing_table = Table(table_name, metadata, autoload_with=engine)

	used_names = {'id'}
	rename_map: dict[str, str] = {}
	for column in existing_table.columns:
		base = 'source_id' if is_id_name(column.name) else column.name
		rename_map[column.name] = unique_column_name(base, used_names)

	temp_table_name = make_temp_table_name(table_name)
	temp_metadata = MetaData()
	temp_columns = [Column('id', Integer(), primary_key=True, autoincrement=True)]

	for column in existing_table.columns:
		target_name = rename_map[column.name]
		temp_columns.append(Column(target_name, column.type, nullable=column.nullable))

	temp_table = Table(temp_table_name, temp_metadata, *temp_columns)

	with engine.begin() as connection:
		inspector = inspect(connection)
		if temp_table_name in inspector.get_table_names():
			Table(temp_table_name, MetaData(), autoload_with=connection).drop(connection)

		temp_table.create(connection)

		select_columns = [
			existing_table.c[source_name].label(target_name)
			for source_name, target_name in rename_map.items()
		]

		if select_columns:
			insert_stmt = temp_table.insert().from_select(
				[target for target in rename_map.values()],
				select(*select_columns),
			)
			connection.execute(insert_stmt)

		existing_table.drop(connection)
		rename_table(connection, connection.dialect.name, temp_table_name, table_name)


def ensure_table(
	engine,
	table_name: str,
	columns: list[str],
	column_types: dict[str, Any],
) -> Table:
	"""Ensure the table exists, has `id` PK, and includes incoming columns."""
	inspector = inspect(engine)
	metadata = MetaData()

	if table_name not in inspector.get_table_names():
		table_columns = []
		table_columns.append(Column('id', Integer(), primary_key=True, autoincrement=True))

		for column in columns:
			sa_type = column_types.get(column, Text())
			table_columns.append(Column(column, sa_type))

		table = Table(table_name, metadata, *table_columns)
		metadata.create_all(engine, tables=[table], checkfirst=True)
		return table

	if not table_has_id_primary_key(engine, table_name):
		rebuild_table_with_id_primary_key(engine, table_name)

	inspector = inspect(engine)
	existing_info = inspector.get_columns(table_name)
	existing_columns = {col['name'] for col in existing_info}
	existing_columns_lower = {str(name).lower() for name in existing_columns}

	preparer = engine.dialect.identifier_preparer
	quoted_table = preparer.quote(table_name)

	with engine.begin() as connection:
		# Add any missing non-id columns with inferred types
		missing_columns = [
			column for column in columns if column.lower() not in existing_columns_lower
		]

		# For existing columns, optionally adjust types on MySQL when inferred type
		# is Text but the DB column is a date/time type (common after previous bad
		# inference). This avoids inserting non-time strings into TIME columns.
		dialect = engine.dialect.name
		if dialect == 'mysql':
			for column in columns:
				if column.lower() not in existing_columns_lower:
					continue
				# find the existing column info
				existing_col = next(
					(c for c in existing_info if str(c['name']).lower() == column.lower()),
					None,
				)
				if not existing_col:
					continue
				sa_type = column_types.get(column, Text())
				try:
					existing_type = existing_col.get('type')
					# compile SQL type names for comparison (fallback to str)
					existing_type_sql = (
						existing_type.compile(engine.dialect)
						if hasattr(existing_type, 'compile')
						else str(existing_type)
					)
					type_sql = sa_type.compile(engine.dialect)
				except Exception:
					existing_type_sql = str(existing_type)
					type_sql = str(sa_type)
				# if inferred is Text but DB has a time/date type, try to modify to TEXT
				if (
					isinstance(sa_type, Text)
					and existing_type_sql
					and any(kw in existing_type_sql.lower() for kw in ('time', 'date'))
				):
					quoted_column = preparer.quote(column)
					try:
						connection.execute(
							text(f'ALTER TABLE {quoted_table} MODIFY COLUMN {quoted_column} TEXT')
						)
					except Exception as exc:  # pragma: no cover - DB-specific behavior
						print(f'Aviso: falha ao alterar tipo da coluna {column}: {exc}')

		for column in missing_columns:
			quoted_column = preparer.quote(column)
			sa_type = column_types.get(column, Text())
			# Try to compile the SQL type name for the dialect; fallback to TEXT
			try:
				type_sql = sa_type.compile(engine.dialect)
			except Exception:
				try:
					type_sql = str(sa_type)
				except Exception:
					type_sql = 'TEXT'

			try:
				connection.execute(
					text(f'ALTER TABLE {quoted_table} ADD COLUMN {quoted_column} {type_sql}')
				)
			except Exception as exc:  # pragma: no cover - DB-specific behavior
				print(f'Aviso: falha ao adicionar coluna {column}: {exc}')

	metadata.reflect(bind=engine, only=[table_name])
	return metadata.tables[table_name]


def import_workbook(engine, file_path: Path) -> int:
	workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
	try:
		worksheet = workbook.active
		rows_iter = worksheet.iter_rows(values_only=True)
		headers_row = next(rows_iter, None)

		if headers_row is None:
			return 0

		# normalize headers (lowercase)
		normalized_headers = normalize_column_names(headers_row)
		columns = remap_incoming_columns(normalized_headers)

		# read rows into memory (sample + payload)
		all_rows = list(rows_iter)

		# infer types using a sample of rows
		projected_rows = [
			tuple(row[idx] if idx < len(row) else None for idx in range(len(columns)))
			for row in all_rows
		]
		filtered_column_types = infer_column_types(columns, projected_rows)

		table = ensure_table(
			engine,
			normalize_table_name(file_path),
			columns,
			filtered_column_types,
		)

		payload = []
		for row in projected_rows:
			if not any(value is not None and value != '' for value in row):
				continue

			row_dict: dict[str, Any] = {}
			for idx, col in enumerate(columns):
				value = row[idx] if idx < len(row) else None
				sa_type = filtered_column_types.get(col, Text())
				row_dict[col] = cast_value_to_type(value, sa_type)

			payload.append(row_dict)

		with engine.begin() as connection:
			if payload:
				connection.execute(table.insert(), payload)

		return len(payload)
	finally:
		workbook.close()


def import_directory(directory: Path, database_url: str) -> dict[str, int]:
	engine = create_engine_from_url(database_url)
	imported_rows: dict[str, int] = {}

	for file_path in sorted(directory.iterdir()):
		if file_path.suffix.lower() not in {'.xlsx', '.xls'} or not file_path.is_file():
			continue
		imported_rows[file_path.name] = import_workbook(engine, file_path)

	return imported_rows


# Argument parsing removed — configure `DATABASE_URL` and `TABLES_DIR` at the top.


def main() -> int:
	tables_dir = TABLES_DIR if isinstance(TABLES_DIR, Path) else Path(TABLES_DIR)

	if not tables_dir.exists():
		raise FileNotFoundError(f'Pasta não encontrada: {tables_dir}')

	imported_rows = import_directory(tables_dir, DATABASE_URL)

	if not imported_rows:
		print('Nenhum arquivo .xlsx foi encontrado para importar.')
		return 0

	for filename, row_count in imported_rows.items():
		print(f'{filename}: {row_count} linhas importadas.')
	return 0


if __name__ == '__main__':
	raise SystemExit(main())
